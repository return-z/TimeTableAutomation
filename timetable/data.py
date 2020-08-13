from openpyxl import load_workbook
import re

wb = load_workbook('timetable.xlsx')
ws = wb['Table 1']
courseCodes = []

for row in ws.iter_rows():
    try:
        if int(row[0].value):
            courseCodes.append(row[1].value)
    except:
        pass

yearWise = {}

for key in range(1,8):
    yearWise[key] = []

for courseCode in courseCodes:  
    code = re.split(r"\s", courseCode)
    yearWise[int(code[1][1])].append(courseCode)

print(yearWise[1])
    

